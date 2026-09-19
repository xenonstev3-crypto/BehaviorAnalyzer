"""Cylinder 分析的图表、Excel、HTML 与 PDF 可复现报告。"""

from __future__ import annotations

import html
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis import CylinderAnalysisResult


def _safe_value(value: Any) -> Any:
    if isinstance(value, (np.floating, np.integer)): return value.item()
    if pd.isna(value): return None
    return value


def _write_frame(sheet, frame: pd.DataFrame, title: str) -> None:
    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=14, color="1F2937")
    if frame.empty:
        sheet.append(["没有记录"]); return
    sheet.append(list(frame.columns))
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in frame.itertuples(index=False, name=None): sheet.append([_safe_value(value) for value in row])
    sheet.freeze_panes = "A3"; sheet.auto_filter.ref = f"A2:{get_column_letter(len(frame.columns))}{len(frame)+2}"
    for index, column in enumerate(frame.columns, start=1):
        sample = [str(column)] + [str(value) for value in frame.iloc[:100, index-1].tolist()]
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(12, max(map(len, sample)) + 2))


def write_excel_report(path: str | Path, result: CylinderAnalysisResult) -> Path:
    """写一个面向复核者的工作簿；不改动原始或 CSV 审计输出。"""
    destination = Path(path)
    if destination.exists(): raise FileExistsError(f"拒绝覆盖已有 Excel：{destination}")
    workbook = Workbook(); summary_sheet = workbook.active; summary_sheet.title = "Summary"; summary_sheet.sheet_view.showGridLines = False
    summary_sheet.append(["Behavior3D Analyzer - Cylinder Summary"]); summary_sheet["A1"].font = Font(bold=True, size=16, color="1F2937")
    summary_sheet.append(["Metric", "Value"])
    for cell in summary_sheet[2]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for key, value in result.summary.items():
        if key in {"config", "created_at_utc"}: continue
        summary_sheet.append([key, _safe_value(value)])
    summary_sheet.append([]); summary_sheet.append(["Configuration", "Value"])
    for cell in summary_sheet[summary_sheet.max_row]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for key, value in result.summary["config"].items(): summary_sheet.append([key, _safe_value(value)])
    summary_sheet.column_dimensions["A"].width = 34; summary_sheet.column_dimensions["B"].width = 48; summary_sheet.freeze_panes = "A3"
    # 可编辑的原生 Excel 柱状图，与摘要统计直接绑定。
    chart_start = 3
    summary_sheet.cell(chart_start, 4, "Event class"); summary_sheet.cell(chart_start, 5, "Count")
    for row, (label, value) in enumerate((("Left", result.summary["left_touch_events"]), ("Right", result.summary["right_touch_events"]), ("Bilateral", result.summary["bilateral_touch_events"])), start=chart_start+1):
        summary_sheet.cell(row, 4, label); summary_sheet.cell(row, 5, value)
    chart = BarChart(); chart.type = "bar"; chart.style = 10; chart.title = "Effective touch events"; chart.y_axis.title = "Event class"; chart.x_axis.title = "Count"; chart.add_data(Reference(summary_sheet, min_col=5, min_row=chart_start, max_row=chart_start+3), titles_from_data=True); chart.set_categories(Reference(summary_sheet, min_col=4, min_row=chart_start+1, max_row=chart_start+3)); chart.height = 7; chart.width = 13; summary_sheet.add_chart(chart, "D8")
    for name, frame, title in (("Events", result.events, "Effective events"), ("Event audit", result.event_audit, "Candidate event audit"), ("Frame quality", result.frame_quality, "Frame-level quality")):
        _write_frame(workbook.create_sheet(name), frame, title)
    workbook.save(destination)
    return destination


def create_plots(output_dir: str | Path, result: CylinderAnalysisResult, *, concise_names: bool = False) -> dict[str, Path]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    directory = Path(output_dir); trial = str(result.summary["trial_id"])
    paths = ({"counts": directory/"counts.png", "timeline": directory/"timeline.png"}
             if concise_names else
             {"counts": directory/f"plot_event_counts_{trial}.png", "timeline": directory/f"plot_event_timeline_{trial}.png"})
    existing = [str(path) for path in paths.values() if path.exists()]
    if existing: raise FileExistsError("拒绝覆盖已有图：" + "；".join(existing))
    counts = [result.summary["left_touch_events"], result.summary["right_touch_events"], result.summary["bilateral_touch_events"]]
    figure, axis = plt.subplots(figsize=(6.4, 3.8)); axis.bar(["Left", "Right", "Bilateral"], counts, color=["#2563EB", "#DC2626", "#7C3AED"]); axis.set_ylabel("Effective events"); axis.set_title("Cylinder effective touch events"); axis.set_ylim(0, max(1, max(counts)+1)); figure.tight_layout(); figure.savefig(paths["counts"], dpi=160); plt.close(figure)
    figure, axis = plt.subplots(figsize=(8, 3.8)); colors = {"left": "#2563EB", "right": "#DC2626", "bilateral": "#7C3AED"}
    for row_index, event in result.events.reset_index(drop=True).iterrows():
        duration = max(float(event["end_time_s"] - event["start_time_s"]), 1 / result.summary["fps"])
        axis.broken_barh([(float(event["start_time_s"]), duration)], (row_index-0.35, 0.7), facecolors=colors[str(event["classification"])])
        axis.text(float(event["start_time_s"]), row_index, str(event["classification"]), va="center", fontsize=8)
    axis.set_xlabel("Time (s)"); axis.set_ylabel("Event"); axis.set_title("Cylinder event timeline"); axis.set_yticks([]); axis.grid(axis="x", alpha=0.25); figure.tight_layout(); figure.savefig(paths["timeline"], dpi=160); plt.close(figure)
    return paths


def _summary_rows(summary: dict[str, Any]) -> str:
    keys = ["trial_id", "animal_id", "coordinate_unit", "frames_analyzed", "left_touch_events", "right_touch_events", "bilateral_touch_events", "total_effective_events", "left_usage_proportion", "right_usage_proportion", "laterality_index", "laterality_formula"]
    return "".join(f"<tr><th>{html.escape(key)}</th><td>{html.escape(str(summary.get(key)))}</td></tr>" for key in keys)


def write_html_report(path: str | Path, result: CylinderAnalysisResult, plots: dict[str, Path]) -> Path:
    destination = Path(path)
    if destination.exists(): raise FileExistsError(f"拒绝覆盖已有 HTML：{destination}")
    events_html = result.events.to_html(index=False, escape=True, classes="events") if not result.events.empty else "<p>没有有效事件。</p>"
    content = f"""<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><title>Cylinder analysis report</title><style>body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:32px;color:#1f2937}}h1,h2{{color:#163b65}}table{{border-collapse:collapse;width:100%;margin:12px 0}}th,td{{border:1px solid #d1d5db;padding:6px;text-align:left;font-size:13px}}th{{background:#e8f0fa}}.events{{font-size:11px;overflow-wrap:anywhere}}img{{max-width:760px;width:100%;border:1px solid #d1d5db;margin:10px 0}}.note{{background:#fff7ed;padding:12px;border-left:4px solid #c2410c}}</style></head><body><h1>Behavior3D Analyzer - Cylinder analysis</h1><p>Preset: {html.escape(result.summary['preset_name'])} ({html.escape(result.summary['preset_version'])})</p><div class='note'>结果依赖重建坐标和所列参数。预设不是唯一国际标准；请结合实验设计复核阈值与事件表。</div><h2>Summary</h2><table>{_summary_rows(result.summary)}</table><h2>Effective event counts</h2><img src='{html.escape(plots['counts'].name)}' alt='Event counts'><h2>Event timeline</h2><img src='{html.escape(plots['timeline'].name)}' alt='Event timeline'><h2>Effective events</h2>{events_html}<h2>Configuration</h2><pre>{html.escape(json.dumps(result.summary['config'], ensure_ascii=False, indent=2))}</pre></body></html>"""
    destination.write_text(content, encoding="utf-8")
    return destination


def write_pdf_report(path: str | Path, result: CylinderAnalysisResult, plots: dict[str, Path]) -> Path:
    """使用 reportlab 的内置 CJK 字体生成可携带的中文 PDF。"""
    destination = Path(path)
    if destination.exists(): raise FileExistsError(f"拒绝覆盖已有 PDF：{destination}")
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        return _write_pdf_report_with_qt(destination, result, plots)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet(); title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName="STSong-Light", fontSize=18, leading=24, textColor=colors.HexColor("#163B65")); body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9, leading=14, alignment=TA_LEFT); heading = ParagraphStyle("HeadingCN", parent=styles["Heading2"], fontName="STSong-Light", fontSize=13, leading=18, textColor=colors.HexColor("#163B65"))
    story = [Paragraph("Behavior3D Analyzer - Cylinder 分析报告", title), Paragraph(f"Trial: {result.summary['trial_id']}　Animal: {result.summary['animal_id']}", body), Paragraph(f"规则：{result.summary['preset_name']}（{result.summary['preset_version']}）", body), Paragraph("说明：结果依赖重建坐标与已保存配置。该预设不是唯一国际标准，应由实验人员复核事件与参数。", body), Spacer(1, 5*mm), Paragraph("统计摘要", heading)]
    summary_keys = ["left_touch_events", "right_touch_events", "bilateral_touch_events", "total_effective_events", "left_usage_proportion", "right_usage_proportion", "laterality_index", "laterality_formula"]
    table = Table([[Paragraph("指标", body), Paragraph("数值", body)]] + [[Paragraph(key, body), Paragraph(str(result.summary[key]), body)] for key in summary_keys], colWidths=[65*mm, 105*mm])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E8F0FA")), ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#D1D5DB")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("FONTNAME", (0,0), (-1,-1), "STSong-Light"), ("PADDING", (0,0), (-1,-1), 5)])); story.extend([table, Spacer(1, 5*mm), Paragraph("有效触壁事件", heading), Image(str(plots["counts"]), width=150*mm, height=89*mm), Paragraph("事件时间轴", heading), Image(str(plots["timeline"]), width=170*mm, height=81*mm)])
    if not result.events.empty:
        rows = [["编号", "分类", "开始帧", "结束帧", "持续秒"]] + [[str(row.get("event_id", "")), str(row.get("classification", "")), str(row.get("start_frame", "")), str(row.get("end_frame", "")), f"{float(row.get('duration_s', 0)):.3f}"] for _, row in result.events.head(30).iterrows()]
        event_table = Table(rows, colWidths=[28*mm, 35*mm, 32*mm, 32*mm, 35*mm]); event_table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E8F0FA")), ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#D1D5DB")), ("FONTNAME", (0,0), (-1,-1), "STSong-Light"), ("PADDING", (0,0), (-1,-1), 4)])); story.extend([Paragraph("逐事件表（最多显示前 30 条；完整记录见 CSV/Excel）", heading), event_table])
    SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm).build(story)
    if not destination.exists() or destination.stat().st_size == 0: raise RuntimeError("PDF 导出失败。")
    return destination


def _write_pdf_report_with_qt(destination: Path, result: CylinderAnalysisResult, plots: dict[str, Path]) -> Path:
    """无 reportlab 时使用已安装的 PySide6 Windows PDF 后端。"""
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QTextDocument
    from PySide6.QtPrintSupport import QPrinter
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    rows = _summary_rows(result.summary)
    events = result.events.head(30).to_html(index=False, escape=True) if not result.events.empty else "<p>没有有效事件。</p>"
    content = f"""<html><head><meta charset='utf-8'><style>body{{font-family:'Microsoft YaHei','Segoe UI',Arial;font-size:9pt;color:#1f2937}}h1,h2{{color:#163b65}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #d1d5db;padding:4px}}th{{background:#e8f0fa}}</style></head><body><h1>Behavior3D Analyzer - Cylinder 分析报告</h1><p>Trial: {html.escape(str(result.summary['trial_id']))}　Animal: {html.escape(str(result.summary['animal_id']))}</p><p>规则：{html.escape(str(result.summary['preset_name']))}（{html.escape(str(result.summary['preset_version']))}）</p><p>结果依赖重建坐标和已保存配置；该预设不是唯一国际标准。</p><h2>统计摘要</h2><table>{rows}</table><h2>有效触壁事件</h2><img src='{plots['counts'].resolve().as_uri()}' width='500'><h2>事件时间轴</h2><img src='{plots['timeline'].resolve().as_uri()}' width='600'><h2>逐事件表</h2>{events}</body></html>"""
    document = QTextDocument(); document.setHtml(content)
    printer = QPrinter(QPrinter.HighResolution); printer.setOutputFormat(QPrinter.PdfFormat); printer.setOutputFileName(str(destination)); printer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Millimeter); document.print_(printer)
    if not destination.exists() or destination.stat().st_size == 0: raise RuntimeError("PySide6 PDF 导出失败。")
    return destination


def write_report_bundle(output_dir: str | Path, result: CylinderAnalysisResult, *, concise_names: bool = False) -> dict[str, Path]:
    """一次生成图、Excel、HTML、PDF；任一同名输出已存在则停止。"""
    directory = Path(output_dir); directory.mkdir(parents=True, exist_ok=True); trial = str(result.summary["trial_id"])
    targets = ({"excel": directory/"report.xlsx", "html": directory/"report.html", "pdf": directory/"report.pdf", "counts_plot": directory/"counts.png", "timeline_plot": directory/"timeline.png"}
               if concise_names else
               {"excel": directory/f"report_{trial}.xlsx", "html": directory/f"report_{trial}.html", "pdf": directory/f"report_{trial}.pdf", "counts_plot": directory/f"plot_event_counts_{trial}.png", "timeline_plot": directory/f"plot_event_timeline_{trial}.png"})
    existing = [str(path) for path in targets.values() if path.exists()]
    if existing: raise FileExistsError("拒绝覆盖已有报告文件：" + "；".join(existing))
    plots = create_plots(directory, result, concise_names=concise_names)
    write_excel_report(targets["excel"], result); write_html_report(targets["html"], result, plots); write_pdf_report(targets["pdf"], result, plots)
    return targets
